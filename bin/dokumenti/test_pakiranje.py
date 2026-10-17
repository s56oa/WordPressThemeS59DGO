"""Pretvorba v .dotx in vpis barvne teme delujeta na pravem zip paketu."""
import zipfile

import docx
import pytest

import pakiranje


@pytest.fixture
def docx_pot(tmp_path):
    pot = tmp_path / "poskus.docx"
    dok = docx.Document()
    dok.add_paragraph("besedilo")
    dok.save(pot)
    return pot


def test_barvna_tema_zapise_barve(docx_pot):
    pakiranje.nastavi_barvno_temo(docx_pot)
    with zipfile.ZipFile(docx_pot) as z:
        tema = z.read("word/theme/theme1.xml").decode("utf-8")
    assert '<a:srgbClr val="15304F"/>' in tema
    assert '<a:srgbClr val="D9702A"/>' in tema
    assert docx.Document(docx_pot).paragraphs[0].text == "besedilo"


def test_v_dotx_zamenja_vrsto_vsebine(docx_pot, tmp_path):
    dotx = pakiranje.v_dotx(docx_pot, tmp_path / "poskus.dotx")
    with zipfile.ZipFile(dotx) as z:
        vrste = z.read("[Content_Types].xml").decode("utf-8")
    assert pakiranje.VRSTA_PREDLOGE in vrste
    assert "wordprocessingml.document.main+xml" not in vrste


def test_dotx_se_se_vedno_odpre(docx_pot, tmp_path):
    dotx = pakiranje.v_dotx(docx_pot, tmp_path / "poskus.dotx")
    assert pakiranje.odpri(dotx).paragraphs[0].text == "besedilo"


def test_izvirni_docx_ostane_nespremenjen(docx_pot, tmp_path):
    pakiranje.v_dotx(docx_pot, tmp_path / "poskus.dotx")
    with zipfile.ZipFile(docx_pot) as z:
        assert "wordprocessingml.document.main+xml" in z.read(
            "[Content_Types].xml"
        ).decode("utf-8")


def test_odpri_deluje_tudi_na_docx(docx_pot):
    assert pakiranje.odpri(docx_pot).paragraphs[0].text == "besedilo"


def test_odpri_ne_spremeni_dotx_na_disku(docx_pot, tmp_path):
    dotx = pakiranje.v_dotx(docx_pot, tmp_path / "poskus.dotx")
    pakiranje.odpri(dotx)
    with zipfile.ZipFile(dotx) as z:
        vrste = z.read("[Content_Types].xml").decode("utf-8")
    assert pakiranje.VRSTA_PREDLOGE in vrste
    assert "wordprocessingml.document.main+xml" not in vrste


def test_nastavi_barvno_temo_zapise_fiksen_cas(docx_pot):
    pakiranje.nastavi_barvno_temo(docx_pot)
    with zipfile.ZipFile(docx_pot) as z:
        for info in z.infolist():
            assert info.date_time == pakiranje.CAS_ZAPISA, info.filename


def test_v_dotx_zapise_fiksen_cas(docx_pot, tmp_path):
    dotx = pakiranje.v_dotx(docx_pot, tmp_path / "poskus.dotx")
    with zipfile.ZipFile(dotx) as z:
        for info in z.infolist():
            assert info.date_time == pakiranje.CAS_ZAPISA, info.filename


def test_ponovno_generiranje_je_bitno_enako(tmp_path):
    """Dva neodvisna teka istega postopka morata dati bitno enak zip paket.

    Posnema pravo generiranje (generiraj.zgradi_vse): dokument se vsakič
    zgradi znova iz kode, ne le ponovno shrani – tako kot pri pravem
    `bin/dokumenti.sh`, kjer mora biti `git status` po ponovnem teku čist.
    """

    def zgradi(pot):
        dok = docx.Document()
        dok.add_paragraph("besedilo")
        dok.save(pot)
        pakiranje.nastavi_barvno_temo(pot)
        return pakiranje.v_dotx(pot, pot.with_suffix(".dotx"))

    prvi = zgradi(tmp_path / "prvi.docx")
    drugi = zgradi(tmp_path / "drugi.docx")

    assert prvi.read_bytes() == drugi.read_bytes()
