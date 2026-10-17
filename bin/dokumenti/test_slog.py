"""Imenovani slogi ustrezajo potrjenemu specu."""
import docx
import pytest
from docx.oxml.ns import qn
from docx.shared import Pt

import slog


@pytest.fixture
def dok():
    d = docx.Document()
    slog.dodaj_sloge(d)
    return d


def test_vsi_slogi_dodani(dok):
    for ime in slog.SLOGI:
        assert ime in [s.name for s in dok.styles]


def test_normal_je_arial_11(dok):
    normal = dok.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size == Pt(11)


def test_telo(dok):
    s = dok.styles["S59DGO telo"]
    assert s.font.size == Pt(11)
    assert s.font.color.rgb == slog.barva("besedilo")
    assert s.paragraph_format.line_spacing == 1.15
    assert s.paragraph_format.space_after == Pt(6)
    assert s.font.bold is False


def test_zadeva_je_krepka_in_modra(dok):
    s = dok.styles["S59DGO zadeva"]
    assert s.font.bold is True
    assert s.font.color.rgb == slog.barva("glava")
    assert s.paragraph_format.space_after == Pt(12)


def test_naslov_je_velik_in_v_velikih_crkah(dok):
    s = dok.styles["S59DGO naslov"]
    assert s.font.size == Pt(14)
    assert s.font.bold is True
    assert s.font.all_caps is True


def test_podnaslov(dok):
    s = dok.styles["S59DGO podnaslov"]
    assert s.font.size == Pt(10)
    assert s.font.all_caps is True
    assert s.font.color.rgb == slog.barva("tiho")


def test_noga_je_drobna(dok):
    s = dok.styles["S59DGO noga"]
    assert s.font.size == slog.VELIKOST_NOGE
    assert str(s.font.color.rgb) == "46586C"


def test_sklep_ima_zamik(dok):
    from docx.shared import Mm

    s = dok.styles["S59DGO sklep"]
    # OOXML hrani mero v celih twipih, zato primerjava v twipih, ne v EMU.
    assert s.paragraph_format.left_indent.twips == Mm(8).twips


def test_meta_ima_desni_tabulator(dok):
    from docx.shared import Mm

    tabulatorji = list(dok.styles["S59DGO meta"].paragraph_format.tab_stops)
    assert len(tabulatorji) == 1
    # OOXML hrani mero v celih twipih, zato primerjava v twipih, ne v EMU.
    assert tabulatorji[0].position.twips == Mm(165).twips


def test_vsi_slogi_uporabljajo_arial(dok):
    for ime in slog.SLOGI:
        assert dok.styles[ime].font.name == "Arial", ime


def test_pisava_arial_nastavi_vse_stiri_atribute():
    dok = docx.Document()
    run = dok.add_paragraph().add_run("besedilo")
    slog.pisava_arial(run)
    rFonts = run._r.rPr.rFonts
    for atribut in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        assert rFonts.get(qn(atribut)) == "Arial"
