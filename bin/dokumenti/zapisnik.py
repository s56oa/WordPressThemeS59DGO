"""Predloga zapisnika seje upravnega odbora; za občni zbor naslovno vrstico uredite ročno."""
from docx.shared import Mm, Pt

import deli
import podatki
import pomocniki
import slog

IME = "S59DGO-zapisnik"

OZNAKE_SEJE = ("Datum", "Kraj", "Začetek", "Konec", "Prisotni", "Opravičeno odsotni")
VZOREC_SEJE = (
    "sobota, 14. marca 2026",
    "klubski prostori, Bazoviška cesta 19A",
    "17.00",
    "19.15",
    "Ognjen Antonič – S56OA, Marko Kranjc, Ana Horvat",
    "en član",
)
VZOREC_TOCK = (
    "Ugotovitev sklepčnosti in potrditev dnevnega reda",
    "Poročilo o delu v letu 2025",
    "Finančno poročilo in načrt za leto 2026",
    "Obnova klubskih naprav na Snežniku",
    "Razno",
)
VZOREC_SKLEPA = "1/1: Upravni odbor potrjuje poročilo o delu za leto 2025."


def zgradi(primer=False):
    dok = deli.nov_dokument()

    naslov = dok.add_paragraph(style="S59DGO naslov")
    # Ista vrednost za primer in za oznako polja: zaporedna številka seje je
    # vedno konkretno število, ne opisno besedilo za izpolnjevanje – zato
    # tudi v predlogi ostane prikazana kot »4.«, ne kot npr. »zaporedna
    # številka seje«.
    stevilka = podatki.polje_ali(primer, "4.", "4.")
    run = naslov.add_run(f"ZAPISNIK {stevilka} SEJE UPRAVNEGA ODBORA")
    slog.pisava_arial(run)
    pomocniki.razmik_znakov(run, 1.0)

    tabela = dok.add_table(rows=len(OZNAKE_SEJE), cols=2)
    pomocniki.brez_obrob(tabela)
    for vrstica, oznaka in enumerate(OZNAKE_SEJE):
        leva = tabela.cell(vrstica, 0)
        leva.width = Mm(42)
        leva.paragraphs[0].style = "S59DGO naslovnik"
        run = leva.paragraphs[0].add_run(oznaka)
        run.font.color.rgb = slog.barva("tiho")
        slog.pisava_arial(run)
        desna = tabela.cell(vrstica, 1)
        desna.width = Mm(123)
        desna.paragraphs[0].style = "S59DGO naslovnik"
        desna.paragraphs[0].add_run(
            podatki.polje_ali(primer, VZOREC_SEJE[vrstica], oznaka.lower())
        )

    dnevni_red = dok.add_paragraph("DNEVNI RED", style="S59DGO podnaslov")
    dnevni_red.paragraph_format.space_before = Pt(18)
    stevilo_tock = len(VZOREC_TOCK) if primer else 3
    tocke = tuple(
        podatki.polje_ali(primer, VZOREC_TOCK[i - 1], f"točka {i}")
        for i in range(1, stevilo_tock + 1)
    )
    for zaporedna, tocka in enumerate(tocke, start=1):
        dok.add_paragraph(f"{zaporedna}. {tocka}", style="S59DGO telo")

    for zaporedna, tocka in enumerate(tocke, start=1):
        naslov_tocke = dok.add_paragraph(style="S59DGO zadeva")
        naslov_tocke.add_run(f"K {zaporedna}. točki: {tocka}")
        dok.add_paragraph(
            podatki.polje_ali(primer, "Razprava in ugotovitve.", "zapis razprave"),
            style="S59DGO telo",
        )
        if zaporedna == 1:
            sklep = dok.add_paragraph(
                f"SKLEP {podatki.polje_ali(primer, VZOREC_SKLEPA, '1/1: besedilo sklepa')}",
                style="S59DGO sklep",
            )
            pomocniki.crta_levo(sklep, podatki.BARVE["poudarek"], 16)

    podpisi = dok.add_table(rows=2, cols=3)
    pomocniki.brez_obrob(podpisi)
    for stolpec, vloga in enumerate(("zapisnikar", "predsedujoči", "overitelj zapisnika")):
        crta = podpisi.cell(0, stolpec).paragraphs[0]
        crta.paragraph_format.space_before = Pt(30)
        # Vrstici s črto in z nazivom vloge naj se ne ločita čez prelom strani.
        crta.paragraph_format.keep_with_next = True
        pomocniki.crta_pod(crta, podatki.BARVE["tiho"], 4)
        vloga_odstavek = podpisi.cell(1, stolpec).paragraphs[0]
        vloga_odstavek.style = "S59DGO naslovnik"
        vloga_odstavek.add_run(vloga)

    return dok
