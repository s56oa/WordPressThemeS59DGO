"""Predloga vabila na dogodek društva."""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

import deli
import podatki
import pomocniki
import slog

IME = "S59DGO-vabilo"

OZNAKE_DOGODKA = ("Kdaj", "Kje", "Dnevni red", "Prijava")
VZOREC_DOGODKA = (
    "sobota, 14. marca 2026, ob 17.00",
    "klubski prostori, Bazoviška cesta 19A, Ilirska Bistrica",
    "poročila, finančni načrt, volitve upravnega odbora",
    "test@s59dgo.si do 10. marca 2026",
)
BELA = RGBColor(0xFF, 0xFF, 0xFF)


def _pocisti(del_dokumenta):
    for element in list(del_dokumenta._element):
        del_dokumenta._element.remove(element)


def _pas(dok, naslov, podnaslov):
    glava = dok.sections[0].first_page_header
    _pocisti(glava)
    tabela = glava.add_table(rows=1, cols=1, width=Mm(200))
    pomocniki.brez_obrob(tabela)
    pomocniki.zamik_tabele(tabela, -20)
    celica = tabela.cell(0, 0)
    celica.width = Mm(200)
    pomocniki.osenci(celica, podatki.BARVE["glava"])

    emblem = celica.paragraphs[0]
    emblem.alignment = WD_ALIGN_PARAGRAPH.CENTER
    emblem.paragraph_format.space_before = Pt(10)
    emblem.add_run().add_picture(str(podatki.VIRI / "logo-klub.png"), width=Mm(28))

    vrstica = celica.add_paragraph()
    vrstica.alignment = WD_ALIGN_PARAGRAPH.CENTER
    vrstica.paragraph_format.space_before = Pt(8)
    run = vrstica.add_run(naslov)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = BELA
    slog.pisava_arial(run)

    pod = celica.add_paragraph()
    pod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pod.add_run(podnaslov)
    run.font.size = Pt(11)
    run.font.color.rgb = slog.barva("pas_besedilo")
    slog.pisava_arial(run)
    pomocniki.razmik_znakov(run, 1.5)

    greben = celica.add_paragraph()
    greben.alignment = WD_ALIGN_PARAGRAPH.CENTER
    greben.paragraph_format.space_before = Pt(8)
    greben.paragraph_format.space_after = Pt(0)
    greben.add_run().add_picture(str(podatki.VIRI / "greben-bel.png"), width=Mm(200))

    glava.add_paragraph()  # Word zahteva odstavek za zadnjo tabelo v glavi


def zgradi(primer=False):
    dok = deli.nov_dokument(z_glavo=False)
    naslov = podatki.polje_ali(primer, "OBČNI ZBOR 2026", "NASLOV DOGODKA")
    _pas(dok, naslov, "Radio klub Snežnik · S59DGO")

    nagovor = "Spoštovani član, vabimo te na redni letni občni zbor društva."
    dok.add_paragraph(
        podatki.polje_ali(primer, nagovor, "nagovor"), style="S59DGO telo"
    )

    tabela = dok.add_table(rows=len(OZNAKE_DOGODKA), cols=2)
    pomocniki.brez_obrob(tabela)
    for vrstica, oznaka in enumerate(OZNAKE_DOGODKA):
        leva = tabela.cell(vrstica, 0)
        leva.width = Mm(38)
        leva.paragraphs[0].style = "S59DGO telo"
        run = leva.paragraphs[0].add_run(oznaka)
        run.bold = True
        run.font.color.rgb = slog.barva("glava")
        slog.pisava_arial(run)
        desna = tabela.cell(vrstica, 1)
        desna.width = Mm(127)
        desna.paragraphs[0].style = "S59DGO telo"
        desna.paragraphs[0].add_run(
            podatki.polje_ali(primer, VZOREC_DOGODKA[vrstica], oznaka.lower())
        )
        pomocniki.crta_pod(desna.paragraphs[0], podatki.BARVE["crta"], 4)

    zakljucek = (
        "Če se dogodka ne moreš udeležiti, nam to sporoči; za potrditev poročil "
        "potrebujemo sklepčnost."
    )
    odstavek = dok.add_paragraph(
        podatki.polje_ali(primer, zakljucek, "zaključno besedilo"), style="S59DGO telo"
    )
    odstavek.paragraph_format.space_before = Pt(12)

    ime = dok.add_paragraph(style="S59DGO telo")
    ime.paragraph_format.space_before = Pt(18)
    ime.paragraph_format.space_after = Pt(0)
    run = ime.add_run(podatki.PODPISNIK["ime"])
    run.bold = True
    slog.pisava_arial(run)
    dok.add_paragraph(podatki.PODPISNIK["funkcija"], style="S59DGO naslovnik")

    return dok
