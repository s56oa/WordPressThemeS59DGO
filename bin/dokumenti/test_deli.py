"""Glava in noga sta na obeh vrstah strani in nosita prave podatke."""
import pytest
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Mm

import deli
import podatki
import slog


@pytest.fixture
def dok():
    return deli.nov_dokument()


def besedilo_dela(del_dokumenta):
    vrstice = [o.text for o in del_dokumenta.paragraphs]
    for tabela in del_dokumenta.tables:
        for vrstica in tabela.rows:
            vrstice += [c.text for c in vrstica.cells]
    return "\n".join(vrstice)


def _odstavki_dela(del_dokumenta):
    """Vsi odstavki dela dokumenta, vključno s tistimi v celicah tabel."""
    odstavki = list(del_dokumenta.paragraphs)
    for tabela in del_dokumenta.tables:
        for vrstica in tabela.rows:
            for celica in vrstica.cells:
                odstavki += celica.paragraphs
    return odstavki


def preveri_arial_v_vseh_runih(del_dokumenta):
    for odstavek in _odstavki_dela(del_dokumenta):
        for run in odstavek.runs:
            rPr = run._r.rPr
            if rPr is None:
                assert run.text == "", "run z besedilom mora imeti nastavljeno pisavo"
                continue
            rFonts = rPr.rFonts
            assert rFonts is not None, "run z besedilom mora imeti nastavljeno pisavo"
            for atribut in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                assert rFonts.get(qn(atribut)) == "Arial"


def test_stran_je_a4_s_specificnimi_robovi(dok):
    r = dok.sections[0]
    # python-docx hrani dolžine v celih twipih, zato mm-vrednosti primerjamo
    # prek .twips namesto z neposrednim ==, ki po zaokrožitvi lahko odpove.
    assert (r.page_width.twips, r.page_height.twips) == (Mm(210).twips, Mm(297).twips)
    assert r.left_margin.twips == Mm(25).twips
    assert r.right_margin.twips == Mm(20).twips
    assert r.top_margin.twips == Mm(42).twips
    assert r.bottom_margin.twips == Mm(25).twips
    assert r.header_distance.twips == Mm(12).twips
    assert r.footer_distance.twips == Mm(8).twips


def test_lezeca_stran():
    d = deli.nov_dokument(lezece=True)
    r = d.sections[0]
    assert r.orientation == WD_ORIENT.LANDSCAPE
    assert r.page_width > r.page_height
    assert r.left_margin.twips == Mm(15).twips


def test_glava_prve_strani_ima_ime_kluba(dok):
    besedilo = besedilo_dela(dok.sections[0].first_page_header)
    assert "RADIO KLUB SNEŽNIK" in besedilo
    assert "ILIRSKA BISTRICA" in besedilo
    assert podatki.KLUB_OZNAKA in besedilo


def test_glava_prve_strani_ima_tri_slike(dok):
    glava = dok.sections[0].first_page_header
    slike = [
        del_ for del_ in glava.part.related_parts.values()
        if del_.content_type.startswith("image/")
    ]
    assert len(slike) == 3, "emblem, ZRS in greben"


def test_glava_nadaljevanja_ima_stevilko_strani(dok):
    besedilo = besedilo_dela(dok.sections[0].header)
    assert "S59DGO" in besedilo
    assert "Stran" in besedilo
    assert "RADIO KLUB SNEŽNIK\nILIRSKA BISTRICA" not in besedilo


def test_noga_je_na_obeh_vrstah_strani(dok):
    for noga in (dok.sections[0].footer, dok.sections[0].first_page_footer):
        besedilo = besedilo_dela(noga)
        assert podatki.DRUSTVO["epota"] in besedilo
        assert podatki.DRUSTVO["splet"] in besedilo
        assert podatki.DRUSTVO["trr"] in besedilo
        assert podatki.DRUSTVO["maticna"] in besedilo
        assert podatki.DRUSTVO["davcna"] in besedilo
        assert "zaporedno številko 21" in besedilo


def test_noga_nima_stare_domene(dok):
    besedilo = besedilo_dela(dok.sections[0].footer)
    assert "s59dgo.org" not in besedilo


def test_noga_ima_tri_stolpce(dok):
    tabela = dok.sections[0].footer.tables[0]
    assert len(tabela.columns) == 3


def test_glava_prve_strani_uporablja_arial_povsod(dok):
    preveri_arial_v_vseh_runih(dok.sections[0].first_page_header)


def test_glava_nadaljevanja_uporablja_arial_povsod(dok):
    preveri_arial_v_vseh_runih(dok.sections[0].header)


def test_noga_register_nima_izrecnega_presledka_po(dok):
    # LibreOffice pri izvozu v PDF izpusti zadnji odstavek noge (vrstico o
    # registru za tabelo), če ima ta izrecen w:after="0"; hrošč se pokaže
    # samo v PDF predogledu (`--pdf`), zato tu preverjamo lastnost odstavka,
    # ne izpisano besedilo, ki bi v testu (brez LibreOffice) izgledalo prav.
    for noga in (dok.sections[0].footer, dok.sections[0].first_page_footer):
        register = noga.paragraphs[-1]
        assert register.text == podatki.DRUSTVO["register"]
        assert register.paragraph_format.space_after is None


def test_noga_uporablja_konstanti_iz_sloga(dok):
    # _odstavek_noge oblikuje neposredno na runu (LibreOffice PDF hrošč, glej
    # njen docstring), a mora uporabiti isti velikosti/barvi kot imenovan
    # slog »S59DGO noga« – prek slog.VELIKOST_NOGE/slog.BARVA_NOGE, ne prek
    # ponovno zapisanih številk, da se vrednosti ne moreta razhajati.
    for noga in (dok.sections[0].footer, dok.sections[0].first_page_footer):
        for odstavek in _odstavki_dela(noga)[1:]:  # [0] je črta nad nogo
            for run in odstavek.runs:
                assert run.font.size == slog.VELIKOST_NOGE
                assert run.font.color.rgb == slog.BARVA_NOGE


def test_crta_noge_je_v_slogu_noge(dok):
    # Brez sloga »S59DGO noga« črta nad nogo deduje 11 pt Normal in je vrstica
    # opazno previsoka (~1,2 mm) na vsaki strani vsakega dokumenta.
    for noga in (dok.sections[0].footer, dok.sections[0].first_page_footer):
        crta = noga.paragraphs[0]
        assert crta.style.name == "S59DGO noga"
