"""Predloga uradnega dopisa društva."""
from docx.shared import Mm, Pt

import deli
import podatki
import pomocniki
import slog

IME = "S59DGO-dopis"

VZOREC_NASLOVNIK = ("Občina Ilirska Bistrica", "Bazoviška cesta 14", "6250 Ilirska Bistrica")
VZOREC_ZADEVA = "Vloga za sofinanciranje obnove klubskih antenskih naprav na Snežniku"
VZOREC_ODSTAVKA = (
    "Radio klub Snežnik Ilirska Bistrica od leta 1964 vzdržuje radioamatersko "
    "dejavnost na območju Ilirske Bistrice in upravlja klubske naprave na Snežniku, "
    "ki so v uporabi tudi ob intervencijah in zvezah v sili.",
    "Antenski sistem je po zimi poškodovan. Za obnovo potrebujemo 4.800 EUR, od tega "
    "društvo zagotavlja polovico iz članarine in prostovoljnega dela. Ponudbe in "
    "fotografije stanja so priložene tej vlogi.",
)


def zgradi(primer=False):
    dok = deli.nov_dokument()

    meta = dok.add_paragraph(style="S59DGO meta")
    datum = podatki.polje_ali(primer, "13. 9. 2026", "datum")
    stevilka = podatki.polje_ali(primer, "04/2026-003", "03/2026-001")
    meta.add_run(f"Ilirska Bistrica, {datum}\tŠt.: {stevilka}")

    naslovnik = (
        VZOREC_NASLOVNIK
        if primer
        else tuple(
            podatki.polje(o)
            for o in ("Naziv prejemnika", "Ulica in hišna številka", "Poštna številka in kraj")
        )
    )
    for zaporedna, vrstica in enumerate(naslovnik):
        odstavek = dok.add_paragraph(style="S59DGO naslovnik")
        run = odstavek.add_run(vrstica)
        if zaporedna == 0:
            run.font.color.rgb = slog.barva("besedilo")
            slog.pisava_arial(run)

    zadeva = dok.add_paragraph(style="S59DGO zadeva")
    oznaka = zadeva.add_run("ZADEVA: ")
    oznaka.font.color.rgb = slog.barva("tiho")
    slog.pisava_arial(oznaka)
    pomocniki.razmik_znakov(oznaka, 0.6)
    zadeva.add_run(podatki.polje_ali(primer, VZOREC_ZADEVA, "Predmet dopisa"))

    if primer:
        odstavki = VZOREC_ODSTAVKA
    else:
        odstavki = (podatki.polje("Prvi odstavek"), podatki.polje("Drugi odstavek"))
    for besedilo in odstavki:
        dok.add_paragraph(besedilo, style="S59DGO telo")

    pozdrav = dok.add_paragraph("Lep pozdrav in 73!", style="S59DGO telo")
    pozdrav.paragraph_format.space_after = Pt(24)

    crta = dok.add_paragraph(style="S59DGO telo")
    crta.paragraph_format.right_indent = Mm(110)
    crta.paragraph_format.space_after = Pt(2)
    pomocniki.crta_pod(crta, podatki.BARVE["tiho"], 4)

    ime = dok.add_paragraph(style="S59DGO telo")
    ime.paragraph_format.space_after = Pt(0)
    run = ime.add_run(podatki.PODPISNIK["ime"])
    run.bold = True
    slog.pisava_arial(run)
    for vrstica in (podatki.PODPISNIK["funkcija"], podatki.PODPISNIK["epota"]):
        odstavek = dok.add_paragraph(vrstica, style="S59DGO naslovnik")
        odstavek.paragraph_format.space_after = Pt(0)

    priloge = dok.add_paragraph(style="S59DGO telo")
    priloge.paragraph_format.space_before = Pt(18)
    run = priloge.add_run("Priloge: ")
    run.bold = True
    slog.pisava_arial(run)
    priloge.add_run(
        podatki.polje_ali(primer, "ponudbi dveh izvajalcev, fotografije stanja", "seznam prilog")
    )

    return dok
