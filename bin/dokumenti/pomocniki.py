"""Nizkonivojski OOXML pripomočki, ki jih python-docx ne ponuja neposredno."""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROBOVI = ("top", "left", "bottom", "right", "insideH", "insideV")


def _element(ime, **atributi):
    el = OxmlElement(f"w:{ime}")
    for kljuc, vrednost in atributi.items():
        el.set(qn(f"w:{kljuc}"), str(vrednost))
    return el


def _tblPr(tabela):
    return tabela._tbl.tblPr


def _pPr(odstavek):
    return odstavek._p.get_or_add_pPr()


def _rPr(run):
    return run._r.get_or_add_rPr()


def _zamenjaj(starsi, ime, nov):
    star = starsi.find(qn(f"w:{ime}"))
    if star is not None:
        starsi.remove(star)
    starsi.append(nov)


def brez_obrob(tabela):
    obrobe = _element("tblBorders")
    for rob in ROBOVI:
        obrobe.append(_element(rob, val="none", sz=0, space=0, color="auto"))
    _zamenjaj(_tblPr(tabela), "tblBorders", obrobe)


def obrobe_tabele(tabela, barva_hex, osmine_pike):
    obrobe = _element("tblBorders")
    for rob in ROBOVI:
        obrobe.append(
            _element(rob, val="single", sz=osmine_pike, space=0, color=barva_hex)
        )
    _zamenjaj(_tblPr(tabela), "tblBorders", obrobe)


def osenci(celica, barva_hex):
    tcPr = celica._tc.get_or_add_tcPr()
    _zamenjaj(tcPr, "shd", _element("shd", val="clear", color="auto", fill=barva_hex))


ROBOVI_ODSTAVKA = ("top", "left", "bottom", "right", "between", "bar")


def _pBdr(odstavek):
    pPr = _pPr(odstavek)
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = _element("pBdr")
        pPr.append(pBdr)
    return pBdr


def _nastavi_rob_odstavka(odstavek, ime_roba, barva_hex, osmine_pike, prostor):
    """Doda/zamenja en rob v obstoječem w:pBdr, ne da bi izbrisal druge robove."""
    pBdr = _pBdr(odstavek)
    obstojeci = pBdr.find(qn(f"w:{ime_roba}"))
    if obstojeci is not None:
        pBdr.remove(obstojeci)
    nov = _element(ime_roba, val="single", sz=osmine_pike, space=prostor, color=barva_hex)
    indeks_novega = ROBOVI_ODSTAVKA.index(ime_roba)
    mesto = len(pBdr)
    for i, otrok in enumerate(pBdr):
        oznaka = otrok.tag.rsplit("}", 1)[-1]
        if oznaka in ROBOVI_ODSTAVKA and ROBOVI_ODSTAVKA.index(oznaka) > indeks_novega:
            mesto = i
            break
    pBdr.insert(mesto, nov)


def crta_pod(odstavek, barva_hex, osmine_pike):
    _nastavi_rob_odstavka(odstavek, "bottom", barva_hex, osmine_pike, 4)


def crta_levo(odstavek, barva_hex, osmine_pike):
    _nastavi_rob_odstavka(odstavek, "left", barva_hex, osmine_pike, 8)


def razmik_znakov(run, pike):
    _zamenjaj(_rPr(run), "spacing", _element("spacing", val=int(round(pike * 20))))


def _polje(odstavek, koda):
    run = odstavek.add_run()
    run._r.append(_element("fldChar", fldCharType="begin"))
    navodilo = OxmlElement("w:instrText")
    navodilo.set(qn("xml:space"), "preserve")
    navodilo.text = f" {koda} "
    run._r.append(navodilo)
    run._r.append(_element("fldChar", fldCharType="separate"))
    run._r.append(_element("fldChar", fldCharType="end"))


def stevilka_strani(odstavek):
    odstavek.add_run("Stran ")
    _polje(odstavek, "PAGE")
    odstavek.add_run(" od ")
    _polje(odstavek, "NUMPAGES")


def zamik_tabele(tabela, mm):
    dvajsetinke = int(round(mm * 56.6929))
    _zamenjaj(_tblPr(tabela), "tblInd", _element("tblInd", w=dvajsetinke, type="dxa"))


def okvir_strani(razdelek, barva_hex, osmine_pike, odmik_pt):
    okvir = _element("pgBorders", offsetFrom="page")
    for rob in ("top", "left", "bottom", "right"):
        okvir.append(
            _element(
                rob,
                val="single",
                sz=osmine_pike,
                space=int(round(odmik_pt)),
                color=barva_hex,
            )
        )
    _zamenjaj(razdelek._sectPr, "pgBorders", okvir)
