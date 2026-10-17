"""Zapisnik ima podatke seje, dnevni red, sklepe in tri podpise."""
from docx.oxml.ns import qn

import podatki
import zapisnik


def test_ime_predloge():
    assert zapisnik.IME == "S59DGO-zapisnik"


def test_naslov_je_v_slogu_naslov():
    dok = zapisnik.zgradi()
    naslovi = [o for o in dok.paragraphs if o.style.name == "S59DGO naslov"]
    assert naslovi and "ZAPISNIK" in naslovi[0].text.upper()


def test_tabela_podatkov_seje():
    tabela = zapisnik.zgradi().tables[0]
    oznake = [tabela.cell(v, 0).text for v in range(len(tabela.rows))]
    for pricakovana in ("Datum", "Kraj", "Začetek", "Konec", "Prisotni", "Opravičeno odsotni"):
        assert pricakovana in oznake, pricakovana


def test_dnevni_red_je_ostevilcen():
    besedilo = "\n".join(o.text for o in zapisnik.zgradi().paragraphs)
    assert "DNEVNI RED" in besedilo
    assert "1. " in besedilo


def test_sklep_ima_oranzno_obrobo():
    dok = zapisnik.zgradi()
    sklepi = [o for o in dok.paragraphs if o.style.name == "S59DGO sklep"]
    assert sklepi, "vsaj en vzorčni sklep"
    obroba = sklepi[0]._p.pPr.find(qn("w:pBdr")).find(qn("w:left"))
    assert obroba.get(qn("w:color")) == podatki.BARVE["poudarek"]
    assert sklepi[0].text.startswith("SKLEP")


def test_trije_podpisi():
    tabela = zapisnik.zgradi().tables[-1]
    assert len(tabela.columns) == 3
    besedilo = " ".join(c.text for c in tabela.rows[-1].cells)
    assert "zapisnikar" in besedilo
    assert "predsedujoči" in besedilo
    assert "overitelj" in besedilo


def test_primer_je_izpolnjen():
    besedilo = "\n".join(o.text for o in zapisnik.zgradi(primer=True).paragraphs)
    assert "»" not in besedilo


def test_podpisna_crta_ostane_z_nazivom_vloge():
    # Brez keep_with_next se vrstica s podpisno črto in vrstica z nazivom
    # vloge pri prelomu strani lahko spet ločita (črta ostane na eni strani,
    # naziv se prestavi na naslednjo).
    tabela = zapisnik.zgradi().tables[-1]
    for stolpec in range(3):
        crta = tabela.cell(0, stolpec).paragraphs[0]
        assert crta.paragraph_format.keep_with_next is True
