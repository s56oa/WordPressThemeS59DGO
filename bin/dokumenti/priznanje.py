"""Predloga priznanja: ležeča stran z dvojnim okvirom, slovensko in angleško."""
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Mm, Pt

import deli
import podatki
import pomocniki
import slog

IME = "S59DGO-priznanje"
IME_EN = "S59DGO-certificate-en"

# Kje na spodnji vrstici se začneta podpisna črta in besedilo funkcije
# (levo od tega ostane prostor za kraj in datum).
ZAMIK_PODPISA = Mm(150)

BESEDILA = {
    "sl": {
        "klub": "RADIO KLUB SNEŽNIK ILIRSKA BISTRICA",
        "naslov": "PRIZNANJE",
        "polje_ime": "ime in priimek – klicni znak",
        "polje_utemeljitev": "utemeljitev",
        "primer_ime": "Janez Novak – S5xABC",
        "primer_utemeljitev": (
            "za uspešno opravljen tečaj za radioamaterja razreda A in izkazano znanje "
            "na izpitu pri Agenciji za komunikacijska omrežja in storitve."
        ),
        "kraj": "Ilirska Bistrica",
        "primer_datum": "14. 3. 2026",
        "polje_datum": "datum",
        "funkcija": podatki.PODPISNIK["funkcija"],
    },
    "en": {
        "klub": "RADIO CLUB SNEŽNIK ILIRSKA BISTRICA",
        "naslov": "CERTIFICATE OF ACHIEVEMENT",
        "polje_ime": "name – callsign",
        "polje_utemeljitev": "reason",
        "primer_ime": "John Smith – G0ABC",
        "primer_utemeljitev": (
            "for outstanding results in the VHF contest organised by Radio Club "
            "Snežnik Ilirska Bistrica."
        ),
        "kraj": "Ilirska Bistrica, Slovenia",
        "primer_datum": "14 March 2026",
        "polje_datum": "date",
        "funkcija": "President of the club",
    },
}


def _sredinski(celica, besedilo, velikost, barva_ime, krepko=False, razmik=None):
    odstavek = celica.add_paragraph()
    odstavek.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = odstavek.add_run(besedilo)
    run.bold = krepko
    run.font.size = Pt(velikost)
    run.font.color.rgb = slog.barva(barva_ime)
    slog.pisava_arial(run)
    if razmik:
        pomocniki.razmik_znakov(run, razmik)
    return odstavek


def _vrstica_podpisa(celica, kraj, datum, funkcija):
    """Spodnja vrstica: levo kraj in datum, desno podpisna črta in funkcija.

    Namesto gnezdene tabele (njeno besedilo `dok.tables` ne bi zajel) sta
    obe polovici en sam odstavek z zamikom/tabulatorjem na `ZAMIK_PODPISA`.
    """
    crta = celica.add_paragraph()
    crta.paragraph_format.space_before = Pt(34)
    crta.paragraph_format.left_indent = ZAMIK_PODPISA
    pomocniki.crta_pod(crta, podatki.BARVE["tiho"], 4)

    besedilo = celica.add_paragraph()
    besedilo.paragraph_format.space_before = Pt(4)
    besedilo.paragraph_format.tab_stops.add_tab_stop(ZAMIK_PODPISA, WD_TAB_ALIGNMENT.LEFT)
    levo = besedilo.add_run(f"{kraj}, {datum}\t")
    levo.font.size = Pt(10)
    slog.pisava_arial(levo)
    desno = besedilo.add_run(funkcija)
    desno.font.size = Pt(10)
    slog.pisava_arial(desno)
    return besedilo


def zgradi(primer=False, jezik="sl"):
    besedila = BESEDILA[jezik]
    dok = deli.nov_dokument(lezece=True)
    pomocniki.okvir_strani(dok.sections[0], podatki.BARVE["glava"], 16, 24)

    okvir = dok.add_table(rows=1, cols=1)
    okvir.alignment = WD_TABLE_ALIGNMENT.CENTER
    pomocniki.obrobe_tabele(okvir, podatki.BARVE["poudarek"], 6)
    sirina_okvira = Mm(255)
    # Nastavljena postavitev (namesto privzetega auto/autofit): brez tega je
    # tblGrid ostal na širini bloka izpred klica (267 mm), medtem ko je bila
    # širina celice spodaj nastavljena na 255 mm – nasprotujoči si meri, ki
    # sta prepuščali dejansko širino notranjega okvira Wordovemu autofitu.
    okvir.autofit = False
    okvir.columns[0].width = sirina_okvira
    celica = okvir.cell(0, 0)
    celica.width = sirina_okvira
    celica.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    celica.paragraphs[0].paragraph_format.space_before = Pt(18)

    logotipa = celica.paragraphs[0]
    logotipa.add_run().add_picture(str(podatki.VIRI / "logo-klub.png"), width=Mm(26))
    logotipa.add_run("    ")
    logotipa.add_run().add_picture(str(podatki.VIRI / "logo-zrs.png"), width=Mm(20))

    _sredinski(celica, besedila["klub"], 10, "tiho", razmik=1.2)
    naslov = _sredinski(celica, besedila["naslov"], 30, "glava", krepko=True, razmik=3.0)
    naslov.paragraph_format.space_before = Pt(14)

    ime = _sredinski(
        celica,
        podatki.polje_ali(primer, besedila["primer_ime"], besedila["polje_ime"]),
        20,
        "glava",
    )
    ime.paragraph_format.space_before = Pt(16)
    ime.paragraph_format.left_indent = Mm(60)
    ime.paragraph_format.right_indent = Mm(60)
    pomocniki.crta_pod(ime, podatki.BARVE["tiho"], 4)

    utemeljitev = _sredinski(
        celica,
        podatki.polje_ali(primer, besedila["primer_utemeljitev"], besedila["polje_utemeljitev"]),
        11,
        "besedilo",
    )
    utemeljitev.paragraph_format.space_before = Pt(10)
    utemeljitev.paragraph_format.left_indent = Mm(45)
    utemeljitev.paragraph_format.right_indent = Mm(45)

    datum = podatki.polje_ali(primer, besedila["primer_datum"], besedila["polje_datum"])
    _vrstica_podpisa(celica, besedila["kraj"], datum, besedila["funkcija"])

    greben = celica.add_paragraph()
    greben.alignment = WD_ALIGN_PARAGRAPH.CENTER
    greben.paragraph_format.space_before = Pt(10)
    greben.add_run().add_picture(str(podatki.VIRI / "greben.png"), width=sirina_okvira)

    return dok
