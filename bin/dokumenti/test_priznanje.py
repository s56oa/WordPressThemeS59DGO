"""Priznanje je ležeče, v dvojnem okviru in obstaja v obeh jezikih."""
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.shared import Mm

import podatki
import priznanje


def test_imeni_predlog():
    assert priznanje.IME == "S59DGO-priznanje"
    assert priznanje.IME_EN == "S59DGO-certificate-en"


def test_lezeca_usmeritev():
    razdelek = priznanje.zgradi().sections[0]
    assert razdelek.orientation == WD_ORIENT.LANDSCAPE
    assert razdelek.page_width > razdelek.page_height


def test_modri_okvir_strani():
    razdelek = priznanje.zgradi().sections[0]
    okvir = razdelek._sectPr.find(qn("w:pgBorders"))
    assert okvir.find(qn("w:top")).get(qn("w:color")) == podatki.BARVE["glava"]


def test_oranzni_notranji_okvir():
    tabela = priznanje.zgradi().tables[0]
    obrobe = tabela._tbl.tblPr.find(qn("w:tblBorders"))
    assert obrobe.find(qn("w:top")).get(qn("w:color")) == podatki.BARVE["poudarek"]


def test_notranji_okvir_ima_stalno_postavitev_in_ujemajoco_sirino():
    # Prej: tblW="auto" in tblGrid pri širini besedilnega bloka (267 mm), a
    # celica na 255 mm – nasprotujoči si širini, prepuščeni Wordovemu
    # autofitu. Zdaj: fiksna postavitev, grid in celica se ujemata.
    tabela = priznanje.zgradi().tables[0]
    tblPr = tabela._tbl.tblPr
    assert tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "fixed"
    gridCol = tabela._tbl.find(qn("w:tblGrid")).find(qn("w:gridCol"))
    sirina_grida = int(gridCol.get(qn("w:w")))
    sirina_celice = tabela.columns[0].cells[0].width.twips
    assert sirina_grida == sirina_celice == Mm(255).twips


def test_slovensko_besedilo():
    besedilo = "\n".join(o.text for t in priznanje.zgradi().tables for v in t.rows for c in v.cells for o in c.paragraphs)
    assert "PRIZNANJE" in besedilo
    assert "RADIO KLUB SNEŽNIK ILIRSKA BISTRICA" in besedilo.upper()
    assert podatki.PODPISNIK["funkcija"] in besedilo


def test_anglesko_besedilo():
    dok = priznanje.zgradi(jezik="en")
    besedilo = "\n".join(o.text for t in dok.tables for v in t.rows for c in v.cells for o in c.paragraphs)
    assert "CERTIFICATE OF ACHIEVEMENT" in besedilo
    assert "President of the club" in besedilo
    assert "PRIZNANJE" not in besedilo


def test_oba_logotipa():
    dok = priznanje.zgradi()
    slike = [d for d in dok.part.related_parts.values() if d.content_type.startswith("image/")]
    assert len(slike) == 3, "klubski emblem, ZRS in greben"


def test_primer_je_izpolnjen():
    dok = priznanje.zgradi(primer=True)
    besedilo = "\n".join(o.text for t in dok.tables for v in t.rows for c in v.cells for o in c.paragraphs)
    assert "»" not in besedilo
