"""Skupni gradniki dokumentov: nastavitev strani, glava in noga."""
import docx
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Mm, Pt

import podatki
import pomocniki
import slog

SIRINA_BESEDILA = Mm(165)


def nastavi_stran(razdelek, lezece=False):
    if lezece:
        razdelek.orientation = WD_ORIENT.LANDSCAPE
        razdelek.page_width, razdelek.page_height = Mm(297), Mm(210)
        for rob in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(razdelek, rob, Mm(15))
    else:
        razdelek.orientation = WD_ORIENT.PORTRAIT
        razdelek.page_width, razdelek.page_height = Mm(210), Mm(297)
        razdelek.left_margin = Mm(25)
        razdelek.right_margin = Mm(20)
        razdelek.top_margin = Mm(42)
        razdelek.bottom_margin = Mm(25)
    razdelek.header_distance = Mm(12)
    razdelek.footer_distance = Mm(8)
    razdelek.different_first_page_header_footer = True


def dodaj_glavo(dok):
    glava = dok.sections[0].first_page_header
    glava.paragraphs[0].text = ""
    tabela = glava.add_table(rows=1, cols=3, width=SIRINA_BESEDILA)
    pomocniki.brez_obrob(tabela)
    for celica, sirina in zip(tabela.rows[0].cells, (Mm(26), Mm(109), Mm(30))):
        celica.width = sirina

    emblem = tabela.cell(0, 0).paragraphs[0]
    emblem.add_run().add_picture(str(podatki.VIRI / "logo-klub.png"), width=Mm(22))

    ime = tabela.cell(0, 1).paragraphs[0]
    ime.paragraph_format.space_after = Pt(0)
    ime.paragraph_format.line_spacing = 1.1
    for zaporedna, vrstica in enumerate(podatki.KLUB_VRSTICI):
        run = ime.add_run(("\n" if zaporedna else "") + vrstica)
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = slog.barva("glava")
        slog.pisava_arial(run)

    oznaka = tabela.cell(0, 1).add_paragraph()
    oznaka.paragraph_format.space_before = Pt(2)
    run = oznaka.add_run(podatki.KLUB_OZNAKA)
    run.font.size = Pt(8)
    run.font.color.rgb = slog.barva("tiho")
    slog.pisava_arial(run)
    pomocniki.razmik_znakov(run, 1.2)

    zrs = tabela.cell(0, 2).paragraphs[0]
    zrs.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    zrs.add_run().add_picture(str(podatki.VIRI / "logo-zrs.png"), width=Mm(18))

    greben = glava.add_paragraph()
    greben.paragraph_format.space_before = Pt(6)
    greben.paragraph_format.space_after = Pt(0)
    greben.add_run().add_picture(str(podatki.VIRI / "greben.png"), width=SIRINA_BESEDILA)
    pomocniki.crta_pod(greben, podatki.BARVE["poudarek"], 8)


def dodaj_glavo_nadaljevanja(dok):
    glava = dok.sections[0].header
    odstavek = glava.paragraphs[0]
    odstavek.paragraph_format.tab_stops.add_tab_stop(SIRINA_BESEDILA, WD_TAB_ALIGNMENT.RIGHT)
    odstavek.add_run("RADIO KLUB SNEŽNIK · S59DGO\t")
    pomocniki.stevilka_strani(odstavek)
    for run in odstavek.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = slog.barva("tiho")
        slog.pisava_arial(run)
    pomocniki.crta_pod(odstavek, podatki.BARVE["crta"], 4)


def _odstavek_noge(odstavek, besedilo):
    """Slog noge nastavi neposredno na runu, ne prek imenovanega sloga.

    LibreOffice je pri izvozu v PDF (`bin/dokumenti.sh --pdf`) znal izpustiti
    zadnji odstavek noge (vrstico o registru) za tabelo s podatki – empirično
    ugotovljeno: pojavi se, če zadnji odstavek uporablja imenovan slog, ali če
    ima izrecno `w:after="0"` v lastnem pPr (glej `_napolni_nogo`, kjer je
    zato pri odstavku o registru ta atribut po klicu odstranjen). V samem
    Wordu ta napaka ni bila opažena; gre torej za posebnost izvoznika za
    predogled, ne za neveljaven OOXML.
    """
    odstavek.paragraph_format.space_after = Pt(0)
    odstavek.paragraph_format.line_spacing = 1.0
    run = odstavek.add_run(besedilo)
    run.font.size = slog.VELIKOST_NOGE
    run.font.color.rgb = slog.BARVA_NOGE
    slog.pisava_arial(run)
    return odstavek


def _napolni_nogo(noga):
    noga.paragraphs[0].text = ""
    d = podatki.DRUSTVO
    stolpci = (
        (d["ime"], d["ulica"], d["posta"]),
        (f"E-pošta: {d['epota']}", f"Splet: {d['splet']}", f"Telefon: {d['telefon']}"),
        (
            f"Matična številka: {d['maticna']}",
            f"Davčna številka: {d['davcna']}",
            f"TRR: {d['trr']}",
        ),
    )
    tabela = noga.add_table(rows=1, cols=3, width=SIRINA_BESEDILA)
    pomocniki.brez_obrob(tabela)
    for stolpec, vrstice in enumerate(stolpci):
        celica = tabela.cell(0, stolpec)
        celica.width = Mm(55)
        pomocniki.osenci(celica, podatki.BARVE["podlaga"])
        for zaporedna, vrstica in enumerate(vrstice):
            odstavek = celica.paragraphs[0] if zaporedna == 0 else celica.add_paragraph()
            _odstavek_noge(odstavek, vrstica)

    register = noga.add_paragraph()
    register.paragraph_format.space_before = Pt(3)
    _odstavek_noge(register, podatki.DRUSTVO["register"])
    register.paragraph_format.space_after = None


def dodaj_nogo(dok):
    razdelek = dok.sections[0]
    crta = razdelek.first_page_footer.paragraphs[0]
    # Brez tega sloga odstavek deduje Footer -> Normal pri 11 pt in je črta
    # nad nogo za ~1,2 mm previsoka; slog »S59DGO noga« (7,5 pt) jo drži tanko.
    crta.style = "S59DGO noga"
    pomocniki.crta_pod(crta, podatki.BARVE["crta"], 4)
    _napolni_nogo(razdelek.first_page_footer)
    crta = razdelek.footer.paragraphs[0]
    crta.style = "S59DGO noga"
    pomocniki.crta_pod(crta, podatki.BARVE["crta"], 4)
    _napolni_nogo(razdelek.footer)


def nov_dokument(lezece=False, z_glavo=True):
    dok = docx.Document()
    nastavi_stran(dok.sections[0], lezece=lezece)
    slog.dodaj_sloge(dok)
    if not lezece:
        if z_glavo:
            dodaj_glavo(dok)
        dodaj_glavo_nadaljevanja(dok)
        dodaj_nogo(dok)
    return dok
