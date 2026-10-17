"""Dopis ima predpisano zaporedje elementov in pravilna polja."""
from docx.shared import Mm

import dopis
import podatki


def test_ime_predloge():
    assert dopis.IME == "S59DGO-dopis"


def test_prazna_predloga_ima_polja():
    besedilo = "\n".join(o.text for o in dopis.zgradi().paragraphs)
    for oznaka in ("Naziv prejemnika", "Ulica in hišna številka", "Poštna številka in kraj", "datum"):
        assert podatki.polje(oznaka) in besedilo, oznaka


def test_zaporedje_slogov():
    slogi = [o.style.name for o in dopis.zgradi().paragraphs if o.text.strip()]
    assert slogi[0] == "S59DGO meta"
    assert slogi[1] == "S59DGO naslovnik"
    assert "S59DGO zadeva" in slogi


def test_zadeva_je_oznacena():
    zadeva = [o for o in dopis.zgradi().paragraphs if o.style.name == "S59DGO zadeva"][0]
    assert zadeva.text.startswith("ZADEVA:")


def test_podpisni_blok():
    besedilo = "\n".join(o.text for o in dopis.zgradi().paragraphs)
    assert podatki.PODPISNIK["ime"] in besedilo
    assert podatki.PODPISNIK["funkcija"] in besedilo
    assert podatki.PODPISNIK["epota"] in besedilo


def test_podpisna_crta_je_siroka_55_mm():
    # python-docx hrani dolžine v celih twipih, zato mm-vrednost primerjamo
    # prek .twips namesto z neposrednim ==, ki po zaokrožitvi lahko odpove.
    crte = [
        o for o in dopis.zgradi().paragraphs
        if o.paragraph_format.right_indent is not None
        and o.paragraph_format.right_indent.twips == Mm(110).twips
    ]
    assert len(crte) == 1


def test_primer_je_izpolnjen():
    besedilo = "\n".join(o.text for o in dopis.zgradi(primer=True).paragraphs)
    assert "»" not in besedilo, "primer ne sme vsebovati praznih polj"
    assert "Občina Ilirska Bistrica" in besedilo


def test_dopis_ima_pozdrav():
    besedilo = "\n".join(o.text for o in dopis.zgradi().paragraphs)
    assert "73" in besedilo, "radioamaterski pozdrav"
