"""Generiranje naredi vseh pet predlog in štiri primere ter jih preveri."""
import zipfile

import docx
import pytest

import generiraj
import pakiranje
import podatki


def zgradi_v(tmp_path):
    predloge = tmp_path / "predloge"
    primeri = tmp_path / "primeri"
    predloge.mkdir()
    primeri.mkdir()
    return generiraj.zgradi_vse(predloge, primeri), predloge, primeri


def _dokument_z_eno_sliko(besedilo=""):
    dok = docx.Document()
    dok.add_paragraph().add_run().add_picture(str(podatki.VIRI / "logo-klub.png"))
    if besedilo:
        dok.add_paragraph(besedilo)
    return dok


def test_nastanejo_vse_datoteke(tmp_path):
    poti, predloge, primeri = zgradi_v(tmp_path)
    imena = sorted(p.name for p in poti)
    assert imena == sorted(
        [
            "S59DGO-dopis.dotx",
            "S59DGO-zapisnik.dotx",
            "S59DGO-vabilo.dotx",
            "S59DGO-priznanje.dotx",
            "S59DGO-certificate-en.dotx",
            "S59DGO-dopis-primer.docx",
            "S59DGO-zapisnik-primer.docx",
            "S59DGO-vabilo-primer.docx",
            "S59DGO-priznanje-primer.docx",
        ]
    )
    assert len(list(predloge.glob("*.docx"))) == 0, "začasne datoteke morajo biti pobrisane"
    assert len(list(primeri.glob("*.dotx"))) == 0


def test_predloge_so_prave_vrste(tmp_path):
    _, predloge, _ = zgradi_v(tmp_path)
    for pot in predloge.glob("*.dotx"):
        with zipfile.ZipFile(pot) as z:
            assert pakiranje.VRSTA_PREDLOGE in z.read("[Content_Types].xml").decode("utf-8")


def test_barvna_tema_je_vpisana(tmp_path):
    _, predloge, _ = zgradi_v(tmp_path)
    with zipfile.ZipFile(predloge / "S59DGO-dopis.dotx") as z:
        tema = z.read("word/theme/theme1.xml").decode("utf-8")
    assert podatki.BARVE["glava"] in tema


def test_predloge_imajo_polja_primeri_pa_ne(tmp_path):
    _, predloge, primeri = zgradi_v(tmp_path)
    predloga = "\n".join(o.text for o in pakiranje.odpri(predloge / "S59DGO-dopis.dotx").paragraphs)
    primer = "\n".join(o.text for o in pakiranje.odpri(primeri / "S59DGO-dopis-primer.docx").paragraphs)
    assert "»" in predloga
    assert "»" not in primer


def test_zgradi_vse_ze_preveri_vsak_izdelek_ob_nastanku(tmp_path):
    # zgradi_vse() od finding #2 dalje kliče preveri() sproti za vsak izdelek,
    # z zastavico ima_nogo iz PREDLOGE (ne več z ugibanjem po imenu datoteke);
    # če pride do sem brez izjeme, je bil vsak izdelek uspešno preverjen.
    poti, _, _ = zgradi_v(tmp_path)
    assert len(poti) == 9


def test_preveri_zazna_staro_domeno(tmp_path):
    pokvarjena = tmp_path / "pokvarjena.docx"
    dok = _dokument_z_eno_sliko("info@s59dgo.org")
    dok.save(pokvarjena)
    with pytest.raises(AssertionError, match="s59dgo.org"):
        generiraj.preveri(pokvarjena, ima_nogo=False, stevilo_slik=1)


def test_preveri_zazna_napacno_stevilo_razdelkov(tmp_path):
    dok = _dokument_z_eno_sliko()
    dok.add_section()
    pot = tmp_path / "dva_razdelka.docx"
    dok.save(pot)
    with pytest.raises(AssertionError, match="razdelek"):
        generiraj.preveri(pot, ima_nogo=False, stevilo_slik=1)


def test_preveri_zazna_napacno_stevilo_slik(tmp_path):
    dok = docx.Document()
    dok.add_paragraph("brez slik")
    pot = tmp_path / "brez_slik.docx"
    dok.save(pot)
    with pytest.raises(AssertionError, match="slik"):
        generiraj.preveri(pot, ima_nogo=False, stevilo_slik=1)


def test_preveri_sprejme_pravo_stevilo_slik(tmp_path):
    dok = _dokument_z_eno_sliko()
    pot = tmp_path / "ena_slika.docx"
    dok.save(pot)
    generiraj.preveri(pot, ima_nogo=False, stevilo_slik=1)  # ne sme sprožiti izjeme


def test_preveri_zazna_premalo_polj_v_predlogi(tmp_path):
    dok = _dokument_z_eno_sliko("»Naziv prejemnika« in »Predmet dopisa« – samo dva para")
    docx_pot = tmp_path / "premalo_polj.docx"
    dok.save(docx_pot)
    dotx_pot = pakiranje.v_dotx(docx_pot, tmp_path / "premalo_polj.dotx")
    with pytest.raises(AssertionError, match="polj"):
        generiraj.preveri(dotx_pot, ima_nogo=False, stevilo_slik=1)


def test_preveri_zavrne_osamljen_odpiralni_narekovaj(tmp_path):
    """Regresija: gola prisotnost »-a ne sme zadostiti preverjanju polj.

    Dva popolna para in en osamljen odpiralni narekovaj brez zaključka:
    naivno preverjanje (`"»" in besedilo`) bi to spregledalo, štetje parov
    z regularnim izrazom pa pravilno zazna, da je polj premalo (< 3).
    """
    dok = _dokument_z_eno_sliko(
        "»Naziv prejemnika« in »Predmet dopisa« ter osamljen » brez zaključka"
    )
    docx_pot = tmp_path / "osamljen_narekovaj.docx"
    dok.save(docx_pot)
    dotx_pot = pakiranje.v_dotx(docx_pot, tmp_path / "osamljen_narekovaj.dotx")
    with pytest.raises(AssertionError, match="polj"):
        generiraj.preveri(dotx_pot, ima_nogo=False, stevilo_slik=1)


def test_preveri_uposteva_zastavico_ima_nogo(tmp_path):
    dok = _dokument_z_eno_sliko()
    pot = tmp_path / "brez_noge.docx"
    dok.save(pot)
    generiraj.preveri(pot, ima_nogo=False, stevilo_slik=1)  # ne sme sprožiti izjeme

    with pytest.raises(AssertionError, match="nogi"):
        generiraj.preveri(pot, ima_nogo=True, stevilo_slik=1)
