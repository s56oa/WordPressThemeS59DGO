"""Vabilo ima modri pas, podatke dogodka in običajno nogo."""
from docx.oxml.ns import qn
from docx.shared import Mm

import podatki
import vabilo


def test_ime_predloge():
    assert vabilo.IME == "S59DGO-vabilo"


def test_pas_je_moder_in_sega_cez_robove():
    glava = vabilo.zgradi().sections[0].first_page_header
    pas = glava.tables[0]
    senca = pas.cell(0, 0)._tc.tcPr.find(qn("w:shd"))
    assert senca.get(qn("w:fill")) == podatki.BARVE["glava"]
    zamik = int(pas._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")))
    assert zamik < 0, "pas mora segati čez levi rob besedila"


def test_pas_ima_emblem_in_beli_greben():
    glava = vabilo.zgradi().sections[0].first_page_header
    slike = [
        d for d in glava.part.related_parts.values() if d.content_type.startswith("image/")
    ]
    assert len(slike) == 2, "emblem in beli greben"


def test_podatki_dogodka():
    tabela = vabilo.zgradi().tables[0]
    oznake = [tabela.cell(v, 0).text for v in range(len(tabela.rows))]
    assert oznake == list(vabilo.OZNAKE_DOGODKA)


def test_noga_ostane_klubska():
    noga = vabilo.zgradi().sections[0].first_page_footer
    besedilo = " ".join(c.text for t in noga.tables for v in t.rows for c in v.cells)
    assert podatki.DRUSTVO["epota"] in besedilo


def test_primer_je_izpolnjen():
    dok = vabilo.zgradi(primer=True)
    besedilo = "\n".join(o.text for o in dok.paragraphs)
    besedilo += " ".join(c.text for t in dok.tables for v in t.rows for c in v.cells)
    assert "»" not in besedilo
    assert "OBČNI ZBOR" in "\n".join(
        o.text for o in dok.sections[0].first_page_header.tables[0].cell(0, 0).paragraphs
    ).upper()
